#!/usr/bin/env python3
"""Parse osnovnaSkolaPitanja.xlsx -> bosanski.json + engleski.json with topic/level tags.

Level/topic mapping from BOS-STAT and ENG-STAT sheets (cells A2-Axx => question #1-...):
Bosanski:
  Književnost: 1-20 OSNOVNI, 21-40 SREDNJI, 41-50 VISI
  Jezik: 51-76 OSNOVNI, 77-102 SREDNJI, 103-115 VISI
  Pravopis: 116-123 OSNOVNI, 124-131 SREDNJI, 132-135 VISI
  Kultura izražavanja: 136-137 OSNOVNI, 138-139 SREDNJI, 140 VISI
  Medijska kultura: 141-144 OSNOVNI, 145-148 SREDNJI, 149-150 VISI
Engleski:
  Vokabular: 1-10 OSNOVNI, 11-20 SREDNJI, 21-30 VISI
  Gramatika: 31-40 OSNOVNI, 41-50 SREDNJI, 51-60 VISI
  Čitanje i razumijevanje: 61-65 OSNOVNI, 66-70 SREDNJI, 71-75 VISI
"""
import json
import re
import openpyxl

SRC = '/home/z/my-project/upload/osnovnaSkolaPitanja.xlsx'
OUT_DIR = '/home/z/my-project/work/quiz/bank'

BOS_MAP = [
    ('Književnost', 1, 20, 'OSNOVNI'), ('Književnost', 21, 40, 'SREDNJI'), ('Književnost', 41, 50, 'VISI'),
    ('Jezik', 51, 76, 'OSNOVNI'), ('Jezik', 77, 102, 'SREDNJI'), ('Jezik', 103, 115, 'VISI'),
    ('Pravopis', 116, 123, 'OSNOVNI'), ('Pravopis', 124, 131, 'SREDNJI'), ('Pravopis', 132, 135, 'VISI'),
    ('Kultura izražavanja', 136, 137, 'OSNOVNI'), ('Kultura izražavanja', 138, 139, 'SREDNJI'), ('Kultura izražavanja', 140, 140, 'VISI'),
    ('Medijska kultura', 141, 144, 'OSNOVNI'), ('Medijska kultura', 145, 148, 'SREDNJI'), ('Medijska kultura', 149, 150, 'VISI'),
]
ENG_MAP = [
    ('Vokabular', 1, 10, 'OSNOVNI'), ('Vokabular', 11, 20, 'SREDNJI'), ('Vokabular', 21, 30, 'VISI'),
    ('Gramatika', 31, 40, 'OSNOVNI'), ('Gramatika', 41, 50, 'SREDNJI'), ('Gramatika', 51, 60, 'VISI'),
    ('Čitanje i razumijevanje', 61, 65, 'OSNOVNI'), ('Čitanje i razumijevanje', 66, 70, 'SREDNJI'), ('Čitanje i razumijevanje', 71, 75, 'VISI'),
]

PREFIX_RE = re.compile(r'^\s*[a-dA-D]\)\s*')
READ_PREFIX = re.compile(r'^\s*Read text and then choose answer between A,\s*B,\s*C or D\s*', re.I)


def clean(s):
    if s is None:
        return ''
    s = str(s).replace('\u00a0', ' ')
    s = re.sub(r'[ \t]+', ' ', s)
    s = re.sub(r' *\n *', '\n', s)
    return s.strip()


def strip_opt(s):
    return clean(PREFIX_RE.sub('', str(s or '')))


def norm(s):
    return re.sub(r'\s+', ' ', clean(s).lower().rstrip('.').strip())


def tag_for(no, mapping):
    for topic, a, b, lvl in mapping:
        if a <= no <= b:
            return topic, lvl
    raise ValueError(f'no mapping for question #{no}')


def parse_sheet(ws, count, mapping, subject):
    out = []
    for r in range(2, count + 2):
        raw_q = ws.cell(row=r, column=1).value
        opts_raw = [ws.cell(row=r, column=c).value for c in range(2, 6)]
        correct_raw = ws.cell(row=r, column=6).value
        q = clean(raw_q)
        opts = [strip_opt(o) for o in opts_raw]
        correct = strip_opt(correct_raw)
        no = r - 1
        # split English reading passages on QUESTION:
        passage = None
        m = re.split(r'\bQUESTION\s*:', q, maxsplit=1, flags=re.I)
        if len(m) == 2:
            passage = READ_PREFIX.sub('', m[0]).strip()
            q = m[1].strip()
        else:
            q2 = READ_PREFIX.sub('', q).strip()
            if q2 != q:
                q = q2  # prefix without QUESTION: split
        # find correct index
        idx = None
        for i, o in enumerate(opts):
            if norm(o) == norm(correct):
                idx = i
                break
        if idx is None:
            # relaxed: startswith match
            for i, o in enumerate(opts):
                if norm(o).startswith(norm(correct)[:30]) or norm(correct).startswith(norm(o)[:30]):
                    idx = i
                    break
        if idx is None:
            print(f'!! {subject} #{no}: correct not matched: {correct!r} vs opts {opts}')
            continue
        topic, level = tag_for(no, mapping)
        out.append({
            'original_no': no,
            'subject': subject,
            'topic': topic,
            'level': level,
            'question': q,
            'passage': passage,
            'options': opts,
            'correct_index': idx,
        })
    return out


def main():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    bos = parse_sheet(wb['BosanskiJezik'], 150, BOS_MAP, 'BOSANSKI')
    eng = parse_sheet(wb['engleski jezik'], 75, ENG_MAP, 'ENGLESKI')
    print(f'bosanski: {len(bos)}, engleski: {len(eng)}')
    with open(f'{OUT_DIR}/bosanski.json', 'w', encoding='utf-8') as f:
        json.dump(bos, f, ensure_ascii=False, indent=1)
    with open(f'{OUT_DIR}/engleski.json', 'w', encoding='utf-8') as f:
        json.dump(eng, f, ensure_ascii=False, indent=1)
    # sample outputs
    print(json.dumps(bos[0], ensure_ascii=False, indent=1)[:400])
    for q in eng:
        if q['passage']:
            print(json.dumps(q, ensure_ascii=False, indent=1)[:600])
            break
    # topics check
    from collections import Counter
    print('BOS topics:', Counter((q['topic'], q['level']) for q in bos))
    print('ENG topics:', Counter((q['topic'], q['level']) for q in eng))


if __name__ == '__main__':
    main()
